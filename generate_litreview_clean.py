#!/usr/bin/env python3
"""Generate formatted Literature Review DOCX - Clean Version"""
import docx
from docx.shared import Pt, RGBColor, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

def add_heading(doc, text, level=1, color=(195, 97, 47)):
    h = doc.add_heading(text, level=level)
    run = h.runs[0]
    run.font.color.rgb = RGBColor(*color)
    run.font.name = 'Georgia'
    return h

def add_para(doc, text, bold=False, italic=False, color=None, highlight=None):
    p = doc.add_paragraph()
    run = p.add_run(text)
    run.font.size = Pt(11)
    run.font.name = 'Calibri'
    if bold: run.bold = True
    if italic: run.italic = True
    if color: run.font.color.rgb = RGBColor(*color)
    if highlight:
        shading = OxmlElement('w:shd')
        shading.set(qn('w:fill'), highlight)
        run._element.get_or_add_rPr().append(shading)
    return p

def add_citation(doc, authors, year, title, venue, arxiv):
    p = doc.add_paragraph(style='List Bullet')
    p.add_run(f"{authors} ({year}). ").bold = True
    p.add_run(f'"{title}". ')
    p.add_run(venue).italic = True
    p.add_run(f". arXiv:{arxiv}")
    p.paragraph_format.left_indent = Inches(0.25)

doc = docx.Document()

# ============ TITLE PAGE ============
title = doc.add_heading('Literature Review', 0)
title.alignment = WD_ALIGN_PARAGRAPH.CENTER
title.runs[0].font.color.rgb = RGBColor(196, 97, 47)
title.runs[0].font.size = Pt(28)

subtitle = doc.add_paragraph('Do Audio Transformers Need Registers?')
subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
subtitle.runs[0].font.size = Pt(16)
subtitle.runs[0].italic = True

meta = doc.add_paragraph('Research Proposal Literature Survey')
meta.alignment = WD_ALIGN_PARAGRAPH.CENTER
meta.runs[0].font.size = Pt(12)
meta.runs[0].font.color.rgb = RGBColor(92, 99, 93)

doc.add_paragraph()
doc.add_page_break()

# ============ SECTION 1: EXECUTIVE SUMMARY ============
add_heading(doc, '1. Executive Summary', level=1)
add_para(doc,
"""This literature review surveys the emerging research area of attention artifacts and register tokens in vision transformers, and identifies a critical research gap: these phenomena have never been systematically studied in audio transformers. We analyze 13 recent papers (2021–2026) spanning the discovery of high-norm artifact tokens in vision models, proposed solutions via register tokens, mechanistic studies of attention sinks, and the few existing audio transformer architectures.""")

add_para(doc, "", highlight='FFEB99')
p = doc.paragraphs[-1]
p.add_run("Key Finding: ").bold = True
p.add_run("Zero papers have analyzed whether pretrained audio encoders (AST, HuBERT, WavLM, Whisper, BEATs) exhibit the register/artifact phenomenon that affects nearly all vision transformers.")

add_para(doc,
"""Our proposed research will transfer the test-time register method from Jiang et al. (NeurIPS 2025 spotlight) to audio transformers, providing the first systematic cross-model study of attention artifacts in the audio domain. This work addresses a verified gap with low compute requirements (~50–80 Kaggle GPU-hours) and clear publication targets (ICASSP 2027, Interspeech 2027).""")

doc.add_page_break()
